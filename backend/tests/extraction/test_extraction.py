"""Golden expectations plus independent python-pptx and exact geometry oracles."""

import hashlib
import io
import json
import os
import subprocess
import sys
from collections import Counter
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from pptx import Presentation

from buddy_api.extraction import ExtractionError, Limits, extract_pptx
from buddy_api.extraction.cli import main
from buddy_api.extraction.contracts import Extraction

ROOT = Path(__file__).resolve().parents[3]
FIXTURES = ROOT / "tests/fixtures/slides"
CONTROL = FIXTURES / "Buddy_Extraction_Geometry.pptx"


def external_fixture(filename, expected):
    directory = os.environ.get("BUDDY_PPTX_FIXTURE_DIR")
    if directory is None:
        pytest.skip("Optional local Drive corpus: set BUDDY_PPTX_FIXTURE_DIR to enable.")
    if not directory.strip():
        pytest.fail("BUDDY_PPTX_FIXTURE_DIR must name a local fixture directory.")
    path = Path(directory) / filename
    if not path.is_file():
        pytest.fail(f"Configured external fixture is missing: {filename}")
    if sha(path) != expected["sha256"]:
        pytest.fail(f"Configured external fixture SHA-256 does not match: {filename}")
    return path


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def rewrite(tmp_path, updates=None, remove=(), source=CONTROL):
    path = tmp_path / "adversarial.pptx"
    with ZipFile(source) as original, ZipFile(path, "w", ZIP_DEFLATED) as result:
        for entry in original.infolist():
            if entry.filename in remove:
                continue
            data = original.read(entry)
            if updates and entry.filename in updates:
                replacement = updates[entry.filename]
                data = replacement(data) if callable(replacement) else replacement
            result.writestr(entry.filename, data)
        for name, data in (updates or {}).items():
            if name not in original.namelist():
                result.writestr(name, data)
    return path


def codes(result):
    return {w.code for w in result.warnings}


def test_known_geometry_and_content():
    result = extract_pptx(CONTROL, source_revision="controlled-v1")
    assert len(result.slides) == 2
    slide = result.slides[0]
    assert (slide.width, slide.height) == (9144000, 12192000)
    assert slide.notes == "Speaker notes for the controlled fixture."
    assert result.slides[1].notes is None
    assert "Paragraph 249: dense extraction content." in result.slides[1].text
    assert Counter(o.kind for o in slide.objects) == {
        "shape": 2,
        "group": 1,
        "table": 1,
        "chart": 1,
        "picture": 1,
    }
    rotated, group, child, table, chart, image = slide.objects
    assert rotated.text == "Rotated title"
    assert rotated.bounds.model_dump() == {
        "x": 1500000,
        "y": 500000,
        "width": 1000000,
        "height": 2000000,
    }
    assert [(p.x, p.y) for p in rotated.corners] == [
        (2500000, 500000),
        (2500000, 2500000),
        (1500000, 2500000),
        (1500000, 500000),
    ]
    assert child.parent_id == group.id
    assert child.bounds == group.bounds
    assert child.bounds.model_dump() == {
        "x": 1100000,
        "y": 1900000,
        "width": 600000,
        "height": 800000,
    }
    assert [(p.x, p.y) for p in child.corners] == [
        (1700000, 2700000),
        (1700000, 1900000),
        (1100000, 1900000),
        (1100000, 2700000),
    ]
    assert table.text == "Name\nValue\nSample\n42"
    assert chart.text == "Measurements\nAlpha\nBeta\n2\n5"
    assert image.bounds.model_dump() == {
        "x": 5000000,
        "y": 1000000,
        "width": 2000000,
        "height": 1000000,
    }
    assert {"chart_cache_only", "notes_unavailable", "embedded_material_not_executed"} <= codes(
        result
    )
    assert Extraction.model_validate_json(result.model_dump_json()) == result


def shape_texts(shapes):
    for shape in shapes:
        if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            yield from shape_texts(shape.shapes)
        elif shape.has_text_frame:
            yield shape.shape_id, shape.text.replace("\v", "\n")
        elif shape.has_table:
            yield shape.shape_id, "\n".join(cell.text for cell in shape.table.iter_cells())


@pytest.mark.parametrize(
    "filename",
    [
        "Buddy_60_Second_Pitch.pptx",
        "Dickinson_Sample_Slides.pptx",
        "SustainABLE_Orientation_presentation.pptx",
        "WWF_Wildlife_and_Climate_Change.pptx",
    ],
)
def test_approved_drive_decks_against_golden_and_independent_reader(filename):
    expected = json.loads((FIXTURES / "golden.json").read_text())[filename]
    path = external_fixture(filename, expected)
    before = sha(path)
    result = extract_pptx(path, source_revision="approved-fixture-v1")
    assert before == sha(path) == result.source_sha256 == expected["sha256"]
    oracle = Presentation(path)
    assert len(result.slides) == len(oracle.slides) == expected["slides"]
    assert sum(len(s.objects) for s in result.slides) == expected["objects"]
    for actual, slide in zip(result.slides, oracle.slides, strict=True):
        assert actual.source_id == slide.slide_id
        assert (actual.width, actual.height) == (oracle.slide_width, oracle.slide_height)
        by_id = {o.source_id: o for o in actual.objects}
        for identifier, text in shape_texts(slide.shapes):
            assert by_id[identifier].text == text
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            assert actual.notes == slide.notes_slide.notes_text_frame.text.replace("\v", "\n")
    assert (
        hashlib.sha256("\n".join(s.text for s in result.slides).encode()).hexdigest()
        == expected["text_sha256"]
    )


def test_external_corpus_is_opt_in_even_if_a_file_exists(tmp_path, monkeypatch):
    monkeypatch.delenv("BUDDY_PPTX_FIXTURE_DIR", raising=False)
    with pytest.raises(pytest.skip.Exception, match="Optional local Drive corpus"):
        external_fixture(CONTROL.name, {"sha256": sha(CONTROL)})


@pytest.mark.parametrize("directory", ["", "missing"])
def test_configured_external_corpus_cannot_silently_skip(directory, tmp_path, monkeypatch):
    monkeypatch.setenv("BUDDY_PPTX_FIXTURE_DIR", "" if not directory else str(tmp_path))
    with pytest.raises(pytest.fail.Exception):
        external_fixture("not-present.pptx", {"sha256": "0" * 64})


def test_external_hash_is_checked_before_parsing(tmp_path, monkeypatch):
    monkeypatch.setenv("BUDDY_PPTX_FIXTURE_DIR", str(tmp_path))
    path = tmp_path / "changed.pptx"
    path.write_bytes(b"not the approved bytes")
    with pytest.raises(pytest.fail.Exception, match="SHA-256 does not match"):
        external_fixture(path.name, {"sha256": "0" * 64})
    assert external_fixture(path.name, {"sha256": sha(path)}) == path


def test_repeated_stream_and_path_results_are_stable_and_revision_scoped():
    first = extract_pptx(CONTROL, source_revision="v1")
    stream = io.BytesIO(CONTROL.read_bytes())
    assert extract_pptx(stream, source_revision="v1") == first
    assert not stream.closed
    second = extract_pptx(CONTROL, source_revision="v2")
    assert [s.id for s in first.slides] != [s.id for s in second.slides]
    assert len({s.id for s in first.slides}) == len(first.slides)


@pytest.mark.parametrize(
    "limits,code",
    [
        ({"source_bytes": 100}, "source_limit"),
        ({"expanded_bytes": 100}, "archive_limit"),
        ({"part_bytes": 100}, "xml_limit"),
        ({"entries": 1}, "archive_limit"),
        ({"xml_nodes": 10}, "xml_limit"),
        ({"xml_depth": 3}, "xml_limit"),
        ({"slides": 1}, "slide_limit"),
        ({"objects": 1}, "object_limit"),
        ({"text_characters": 10}, "text_limit"),
        ({"result_bytes": 10}, "result_limit"),
        ({"seconds": 0.05}, "time_limit"),
    ],
)
def test_limits_fail_closed(limits, code):
    with pytest.raises(ExtractionError) as exc:
        extract_pptx(CONTROL, source_revision="v1", limits=Limits(**limits))
    assert exc.value.code == code


@pytest.mark.parametrize("revision", ["", "private/path", "x" * 129, "bad\nrevision"])
def test_invalid_revision(revision):
    with pytest.raises(ExtractionError, match="invalid_source_revision"):
        extract_pptx(CONTROL, source_revision=revision)


@pytest.mark.parametrize("payload", [b"", b"secret private content", b"PK\x03\x04broken"])
def test_invalid_archive_is_sanitized(payload):
    with pytest.raises(ExtractionError) as exc:
        extract_pptx(io.BytesIO(payload), source_revision="v1")
    assert "secret" not in str(exc.value)
    assert "private" not in str(exc.value)


@pytest.mark.parametrize(
    "name", ["../outside.xml", "/outside.xml", "ppt\\evil.xml", "C:/outside.xml", "ppt/./evil.xml"]
)
def test_traversal_entries_are_never_written(tmp_path, name):
    path = rewrite(tmp_path, {name: b"<x/>"})
    if "\\" in name:
        # Windows ZipInfo normalizes separators on write; patch both directory entries.
        path.write_bytes(path.read_bytes().replace(name.replace("\\", "/").encode(), name.encode()))
    with pytest.raises(ExtractionError, match="unsafe_archive"):
        extract_pptx(path, source_revision="v1")
    assert not (tmp_path / "outside.xml").exists()


def test_memory_ceiling_is_enforced_by_the_os():
    from buddy_api.extraction.runner import parser_memory_limit

    script = (
        "import sys; sys.stdin.buffer.read(1); "
        "from buddy_api.extraction.worker import set_posix_limits; "
        "from buddy_api.extraction import Limits; "
        "set_posix_limits(Limits(memory_mb=128)); "
        "exec('import time\\nblocks=[]\\nfor _ in range(32):\\n "
        "blocks.append(bytearray(8*1024*1024)); time.sleep(0.02)')"
    )
    process = subprocess.Popen(
        [sys.executable, "-c", script],
        stdin=subprocess.PIPE,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    try:
        if sys.platform == "darwin":
            with (
                pytest.raises(ExtractionError, match="memory_limit"),
                parser_memory_limit(process.pid, 128 * 1024 * 1024),
            ):
                process.communicate(b"x", timeout=15)
        else:
            with parser_memory_limit(process.pid, 128 * 1024 * 1024):
                _, stderr = process.communicate(b"x", timeout=15)
                assert process.returncode != 0
                assert b"MemoryError" in stderr
    finally:
        if process.poll() is None:
            process.kill()
        process.communicate()


def test_rss_watchdog_kills_and_reports_excess_memory():
    from buddy_api.extraction.runner import rss_memory_limit

    process = subprocess.Popen(
        [sys.executable, "-c", "import time; time.sleep(15)"],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    try:
        with pytest.raises(ExtractionError, match="memory_limit"), rss_memory_limit(process.pid, 1):
            process.communicate(timeout=10)
        assert process.returncode != 0
    finally:
        if process.poll() is None:
            process.kill()
        process.communicate()


def test_machine_readable_schema_matches_contract():
    assert (
        json.loads((ROOT / "docs/extraction.schema.json").read_text())
        == Extraction.model_json_schema()
    )


@pytest.mark.parametrize(
    "xml",
    [
        b'<!DOCTYPE x [<!ENTITY a "private-secret">]><x>&a;</x>',
        b'<!DOCTYPE x SYSTEM "file:///private-secret"><x/>',
        b"<x><unclosed>",
    ],
)
def test_unused_xml_cannot_hide_dtd_entities_or_malformed_content(tmp_path, xml):
    path = rewrite(tmp_path, {"ppt/unused.xml": xml})
    with pytest.raises(ExtractionError, match="invalid_package"):
        extract_pptx(path, source_revision="v1")


def test_missing_part_and_duplicate_ids_are_not_silently_omitted(tmp_path):
    path = rewrite(tmp_path, remove=("ppt/slides/slide1.xml",))
    with pytest.raises(ExtractionError, match="missing_part"):
        extract_pptx(path, source_revision="v1")
    path = rewrite(
        tmp_path, {"ppt/presentation.xml": lambda b: b.replace(b'id="257"', b'id="256"')}
    )
    with pytest.raises(ExtractionError, match="duplicate_slide_id"):
        extract_pptx(path, source_revision="v1")


def test_external_links_are_ignored_without_network(tmp_path):
    addition = (
        b'<Relationship Id="external" '
        b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" '
        b'Target="http://127.0.0.1:1/never-request" TargetMode="External"/>'
    )
    path = rewrite(
        tmp_path,
        {
            "ppt/slides/_rels/slide1.xml.rels": lambda b: b.replace(
                b"</Relationships>", addition + b"</Relationships>"
            )
        },
    )
    assert "external_links_not_followed" in codes(extract_pptx(path, source_revision="v1"))


@pytest.mark.parametrize(
    "target",
    [
        "../../../outside.xml",
        "https://example.com/evil.xml",
        "file:///private.xml",
        "..%2f..%2f..%2foutside.xml",
    ],
)
def test_unsafe_internal_relationships_fail(tmp_path, target):
    path = rewrite(
        tmp_path,
        {
            "ppt/_rels/presentation.xml.rels": lambda b: b.replace(
                b'Target="slides/slide1.xml"', ('Target="' + target + '"').encode()
            )
        },
    )
    with pytest.raises(ExtractionError):
        extract_pptx(path, source_revision="v1")


def test_duplicate_archive_members_fail(tmp_path):
    path = rewrite(tmp_path)
    with pytest.warns(UserWarning), ZipFile(path, "a") as archive:
        archive.writestr("ppt/presentation.xml", b"<private/>")
    with pytest.raises(ExtractionError, match="unsafe_archive"):
        extract_pptx(path, source_revision="v1")


def test_unsupported_main_content_type_fails(tmp_path):
    path = rewrite(
        tmp_path,
        {
            "[Content_Types].xml": lambda b: b.replace(
                b"presentationml.presentation.main+xml", b"bogus+xml"
            )
        },
    )
    with pytest.raises(ExtractionError, match="unsupported_presentation"):
        extract_pptx(path, source_revision="v1")


def test_missing_shape_geometry_is_explicit_and_group_geometry_is_required(tmp_path):
    import re

    path = rewrite(
        tmp_path,
        {"ppt/slides/slide1.xml": lambda b: re.sub(b"<a:xfrm[^>]*>.*?</a:xfrm>", b"", b, count=1)},
    )
    result = extract_pptx(path, source_revision="v1")
    assert result.slides[0].objects[0].bounds is None
    assert "missing_geometry" in codes(result)
    path = rewrite(
        tmp_path,
        {
            "ppt/slides/slide1.xml": lambda b: re.sub(
                b"<p:grpSpPr>.*?</p:grpSpPr>", b"<p:grpSpPr/>", b
            )
        },
    )
    with pytest.raises(ExtractionError, match="invalid_group_geometry"):
        extract_pptx(path, source_revision="v1")
    path = rewrite(
        tmp_path, {"ppt/slides/slide1.xml": lambda b: re.sub(b"<a:chOff[^>]*/>", b"", b)}
    )
    with pytest.raises(ExtractionError, match="invalid_group_geometry"):
        extract_pptx(path, source_revision="v1")


def test_negative_or_nonfinite_dimensions_fail(tmp_path):
    for value in [b"-1", b"NaN", b"Infinity", b"9" * 100]:
        path = rewrite(
            tmp_path,
            {
                "ppt/presentation.xml": lambda b, value=value: b.replace(
                    b'cx="9144000"', b'cx="' + value + b'"'
                )
            },
        )
        with pytest.raises(ExtractionError, match="invalid_geometry"):
            extract_pptx(path, source_revision="v1")


def test_nested_groups_and_inherited_placeholder_geometry():
    from xml.etree.ElementTree import fromstring

    from buddy_api.extraction.geometry import IDENTITY, compose, envelope, point, transform
    from buddy_api.extraction.package import NS, Package
    from buddy_api.extraction.reader import Reader

    def xfrm(off, ext, child_off=(0, 0), child_ext=None):
        child = (
            ""
            if child_ext is None
            else (
                f'<a:chOff x="{child_off[0]}" y="{child_off[1]}"/>'
                f'<a:chExt cx="{child_ext[0]}" cy="{child_ext[1]}"/>'
            )
        )
        return fromstring(
            f'<a:xfrm xmlns:a="{NS["a"]}"><a:off x="{off[0]}" y="{off[1]}"/>'
            f'<a:ext cx="{ext[0]}" cy="{ext[1]}"/>{child}</a:xfrm>'
        )

    _, outer, _, _ = transform(xfrm((100, 200), (800, 600), (10, 20), (400, 200)), IDENTITY)
    _, inner, _, _ = transform(xfrm((50, 60), (80, 60), (0, 0), (40, 20)), outer)
    world, _, w, h = transform(xfrm((2, 3), (10, 5)), inner)
    corners = (point(world, 0, 0), point(world, w, 0), point(world, w, h), point(world, 0, h))
    assert envelope(corners).model_dump() == {"x": 188, "y": 347, "width": 40, "height": 45}
    assert compose(IDENTITY, world) == world
    with Package(str(CONTROL), Limits()) as package:
        reader = Reader(package, "v1", "a" * 64)
        namespaces = f'xmlns:p="{NS["p"]}" xmlns:a="{NS["a"]}"'
        shape = fromstring(
            f'<p:sp {namespaces}><p:nvSpPr><p:nvPr><p:ph idx="3"/></p:nvPr></p:nvSpPr></p:sp>'
        )
        layout = fromstring(
            f"<p:sldLayout {namespaces}><p:cSld><p:spTree><p:sp>"
            '<p:nvSpPr><p:nvPr><p:ph idx="3" type="body"/></p:nvPr></p:nvSpPr>'
            "</p:sp></p:spTree></p:cSld></p:sldLayout>"
        )
        master = fromstring(
            f"<p:sldMaster {namespaces}><p:cSld><p:spTree><p:sp>"
            '<p:nvSpPr><p:nvPr><p:ph type="body"/></p:nvPr></p:nvSpPr>'
            '<p:spPr><a:xfrm><a:off x="10" y="20"/>'
            '<a:ext cx="30" cy="40"/></a:xfrm></p:spPr>'
            "</p:sp></p:spTree></p:cSld></p:sldMaster>"
        )
        inherited = reader.xfrm(shape, [layout, master])
        assert inherited is not None
        assert inherited.find("a:off", NS).get("x") == "10"
        layout_ph = layout.find(".//p:ph", NS)
        master_ph = master.find(".//p:ph", NS)
        for layout_type, master_type in [
            ("obj", "body"),
            ("pic", "body"),
            ("subTitle", "body"),
            ("ctrTitle", "title"),
        ]:
            layout_ph.set("type", layout_type)
            master_ph.set("type", master_type)
            assert reader.xfrm(shape, [layout, master]) is not None


def test_cli_real_invocation_output_and_failure_redaction(tmp_path, capsys):
    output = tmp_path / "extraction.json"
    proc = subprocess.run(
        [
            sys.executable,
            "-m",
            "buddy_api.extraction.cli",
            str(CONTROL),
            "--source-revision",
            "cli-v1",
            "--output",
            str(output),
        ],
        capture_output=True,
        text=True,
        timeout=30,
    )
    assert proc.returncode == 0, proc.stderr
    assert json.loads(proc.stdout)["slides"] == 2
    assert "Grouped body" not in proc.stdout + proc.stderr
    assert Extraction.model_validate_json(output.read_text()).source_revision == "cli-v1"
    before = output.read_bytes()
    assert main([str(CONTROL), "--source-revision", "cli-v1", "--output", str(output)]) == 1
    assert output.read_bytes() == before
    assert "Grouped body" not in capsys.readouterr().err


def test_parser_temporary_files_cleaned_after_success_failure_and_timeout(tmp_path, monkeypatch):
    import tempfile

    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
    extract_pptx(CONTROL, source_revision="v1")
    with pytest.raises(ExtractionError):
        extract_pptx(io.BytesIO(b"bad"), source_revision="v1")
    with pytest.raises(ExtractionError):
        extract_pptx(CONTROL, source_revision="v1", limits=Limits(seconds=0.05))
    assert list(tmp_path.iterdir()) == []
