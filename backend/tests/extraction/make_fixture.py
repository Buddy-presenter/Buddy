"""Rebuild Buddy-authored geometry fixture; deliberately not a production parser."""

import io
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE

FIXTURES = Path(__file__).resolve().parents[3] / "tests/fixtures/slides"


def build(destination: Path) -> None:
    deck = Presentation()
    deck.slide_width, deck.slide_height = 9144000, 12192000
    first = deck.slides.add_slide(deck.slide_layouts[6])
    shape = first.shapes.add_textbox(1000000, 1000000, 2000000, 1000000)
    shape.text = "Rotated title"
    shape.rotation = 90
    group = first.shapes.add_group_shape()
    child = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, 100000, 200000, 400000, 200000)
    child.text = "Grouped body"
    xfrm = group._element.grpSpPr.xfrm
    xfrm.off.x, xfrm.off.y = 1000000, 2000000
    xfrm.ext.cx, xfrm.ext.cy = 800000, 600000
    group.rotation = 90
    xfrm.set("flipH", "1")
    table = first.shapes.add_table(2, 2, 1000000, 4000000, 2000000, 1000000).table
    for cell, text in zip(table.iter_cells(), ["Name", "Value", "Sample", "42"], strict=True):
        cell.text = text
    data = CategoryChartData()
    data.categories = ["Alpha", "Beta"]
    data.add_series("Measurements", [2, 5])
    first.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 1000000, 6000000, 3000000, 2000000, data)
    image = io.BytesIO()
    Image.new("RGB", (32, 16), color=(40, 130, 80)).save(image, format="PNG")
    image.seek(0)
    first.shapes.add_picture(image, 5000000, 1000000, 2000000, 1000000)
    first.notes_slide.notes_text_frame.text = "Speaker notes for the controlled fixture."
    second = deck.slides.add_slide(deck.slide_layouts[6])
    dense = second.shapes.add_textbox(500000, 500000, 8000000, 10000000)
    dense.text = "\n".join(f"Paragraph {i}: dense extraction content." for i in range(250))
    deck.save(destination)


if __name__ == "__main__":
    build(FIXTURES / "Buddy_Extraction_Geometry.pptx")
