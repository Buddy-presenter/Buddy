"""Hostile upload inputs and the real S3 SDK's wire contract, without AWS access."""

import asyncio
import base64
import hashlib
import io
import threading
import zipfile
from datetime import UTC, datetime
from types import SimpleNamespace
from uuid import uuid4

import boto3
import pytest
from botocore.stub import ANY, Stubber
from pptx import Presentation
from pydantic import ValidationError

from buddy_api.app import create_app
from buddy_api.config import ConfigurationError, Settings
from buddy_api.uploads import storage as storage_module
from buddy_api.uploads.config import UploadSettings
from buddy_api.uploads.contracts import PPTX_TYPE, UploadError, UploadMetadata
from buddy_api.uploads.service import blocking
from buddy_api.uploads.storage import SourceStorage
from buddy_api.uploads.validation import validate_source

CONFIG = UploadSettings(
    bucket="buddy-test-private",
    region="eu-central-1",
    kms_key="arn:aws:kms:eu-central-1:123456789012:key/11111111-1111-1111-1111-111111111111",
)


@pytest.fixture
def source(tmp_path):
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    path = tmp_path / "original.pptx"
    presentation.save(path)
    return path


def rewrite(source, additions=(), replace=None):
    original = source.read_bytes()
    with (
        zipfile.ZipFile(io.BytesIO(original)) as old,
        zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as new,
    ):
        for item in old.infolist():
            data = old.read(item)
            if replace:
                data = replace(item.filename, data)
            new.writestr(item, data)
        for name, data in additions:
            new.writestr(name, data)


def test_valid_original_is_not_rewritten(source):
    before = source.read_bytes()
    assert validate_source(source, str(uuid4()), CONFIG) == 1
    assert source.read_bytes() == before


@pytest.mark.parametrize(
    "name",
    [
        "../escape.bin",
        "/absolute.bin",
        "ppt\\escape.bin",
        "C:/escape.bin",
        "ppt/vbaProject.bin",
        "ppt/activeX/control.xml",
        "ppt/embeddings/workbook.xlsx",
        "ppt/media/unsafe.exe",
    ],
)
def test_unsafe_paths_macros_and_embedded_content_are_rejected(source, name):
    rewrite(source, [(name, b"synthetic payload")])
    if "\\" in name:
        # ZipInfo normalizes the native separator when generating on Windows.
        source.write_bytes(
            source.read_bytes().replace(name.replace("\\", "/").encode(), name.encode())
        )
    with pytest.raises(UploadError) as error:
        validate_source(source, str(uuid4()), CONFIG)
    assert error.value.status == 422
    assert name not in error.value.message


def test_duplicate_entries_are_rejected(source):
    with pytest.warns(UserWarning, match="Duplicate name"):
        rewrite(source, [("[Content_Types].xml", b"duplicate")])
    with pytest.raises(UploadError):
        validate_source(source, str(uuid4()), CONFIG)


def test_unused_xml_entities_are_rejected(source):
    rewrite(source, [("unused.xml", b'<!DOCTYPE a [<!ENTITY x "private">]><a>&x;</a>')])
    with pytest.raises(UploadError):
        validate_source(source, str(uuid4()), CONFIG)


def test_macro_content_type_in_utf16_cannot_bypass_checks(source):
    def replace(name, data):
        if name == "[Content_Types].xml":
            text = data.decode().replace(
                "</Types>",
                '<Override PartName="/ppt/hidden.bin" '
                'ContentType="application/vnd.ms-office.vbaProject"/></Types>',
            )
            return text.replace("encoding='UTF-8'", "encoding='UTF-16'").encode("utf-16")
        return data

    rewrite(source, [("ppt/hidden.bin", b"synthetic")], replace)
    with pytest.raises(UploadError):
        validate_source(source, str(uuid4()), CONFIG)


@pytest.mark.parametrize("kind", ["truncated", "encrypted", "expanded", "slides", "compressed"])
def test_corruption_encryption_and_configured_limits(source, kind):
    settings = CONFIG
    if kind == "truncated":
        source.write_bytes(source.read_bytes()[:100])
    elif kind == "encrypted":
        data = bytearray(source.read_bytes())
        for signature, offset in ((b"PK\x03\x04", 6), (b"PK\x01\x02", 8)):
            position = data.index(signature)
            data[position + offset] |= 1
        source.write_bytes(data)
    elif kind == "expanded":
        rewrite(source, [("ppt/media/bomb.bin", b"x" * 1024 * 1024)])
        settings = CONFIG.model_copy(update={"expanded_bytes": 512 * 1024})
    elif kind == "slides":
        presentation = Presentation(str(source))
        presentation.slides.add_slide(presentation.slide_layouts[6])
        presentation.save(source)
        settings = CONFIG.model_copy(update={"max_slides": 1})
    else:
        settings = CONFIG.model_copy(update={"max_bytes": 10})
    with pytest.raises(UploadError):
        validate_source(source, str(uuid4()), settings)


def test_s3_put_uses_checksum_kms_versioning_and_conditional_original(source):
    client = boto3.client(
        "s3",
        region_name="eu-central-1",
        aws_access_key_id="testing",
        aws_secret_access_key="testing",
    )
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    with Stubber(client) as stub:
        stub.add_response(
            "put_object",
            {"VersionId": "immutable-version"},
            {
                "Bucket": CONFIG.bucket,
                "ExpectedBucketOwner": "123456789012",
                "Key": "sources/private/original.pptx",
                "Body": ANY,
                "ContentLength": source.stat().st_size,
                "ContentType": PPTX_TYPE,
                "ChecksumSHA256": base64.b64encode(bytes.fromhex(digest)).decode(),
                "ServerSideEncryption": "aws:kms",
                "SSEKMSKeyId": CONFIG.kms_key,
                "IfNoneMatch": "*",
                "Metadata": {"source-revision": "v1", "sha256": digest},
            },
        )
        storage = SourceStorage(CONFIG, client)
        assert (
            storage.put(
                "sources/private/original.pptx",
                source,
                size=source.stat().st_size,
                sha256=digest,
                revision="v1",
            )
            == "immutable-version"
        )
        stub.assert_no_pending_responses()
    client.close()


def test_s3_cleanup_deletes_versions_and_markers_but_not_a_neighbor():
    client = boto3.client(
        "s3",
        region_name="eu-central-1",
        aws_access_key_id="testing",
        aws_secret_access_key="testing",
    )
    key = "sources/private/original.pptx"
    arguments = {"Bucket": CONFIG.bucket, "ExpectedBucketOwner": "123456789012"}
    listing = {"Prefix": key, "MaxKeys": 100, **arguments}
    with Stubber(client) as stub:
        stub.add_response(
            "list_object_versions",
            {
                "Versions": [{"Key": key, "VersionId": "v1", "LastModified": datetime.now(UTC)}],
                "DeleteMarkers": [{"Key": key, "VersionId": "deleted"}],
                "IsTruncated": False,
            },
            listing,
        )
        for version in ("v1", "deleted"):
            stub.add_response("delete_object", {}, {**arguments, "Key": key, "VersionId": version})
        stub.add_response(
            "list_object_versions",
            {"Versions": [{"Key": key + "-neighbor", "VersionId": "keep"}], "IsTruncated": False},
            listing,
        )
        SourceStorage(CONFIG, client).remove(key)
        stub.assert_no_pending_responses()
    client.close()


def test_cancellation_waits_for_inflight_work_before_cleanup():
    started, finish = threading.Event(), threading.Event()

    def work():
        started.set()
        finish.wait(2)
        return "done"

    async def scenario():
        task = asyncio.create_task(blocking(work))
        while not started.is_set():
            await asyncio.sleep(0)
        task.cancel()
        await asyncio.sleep(0)
        assert not task.done()
        finish.set()
        with pytest.raises(asyncio.CancelledError):
            await task

    asyncio.run(scenario())


def test_cleanup_deadline_defers_remaining_versions_without_unbounded_calls(monkeypatch):
    clock = iter([0, 0, 31])
    monkeypatch.setattr(storage_module, "time", SimpleNamespace(monotonic=lambda: next(clock)))
    client = boto3.client(
        "s3",
        region_name="eu-central-1",
        aws_access_key_id="testing",
        aws_secret_access_key="testing",
    )
    key = "sources/private/original.pptx"
    with Stubber(client) as stub:
        stub.add_response(
            "list_object_versions",
            {
                "Versions": [{"Key": key, "VersionId": "remaining"}],
                "IsTruncated": False,
            },
            {
                "Bucket": CONFIG.bucket,
                "ExpectedBucketOwner": "123456789012",
                "Prefix": key,
                "MaxKeys": 100,
            },
        )
        with pytest.raises(UploadError) as error:
            SourceStorage(CONFIG, client).remove(key)
        assert error.value.code == "cleanup_pending"
        stub.assert_no_pending_responses()
    client.close()


def test_settings_fail_closed_and_disabled_uploads_do_not_need_credentials():
    assert Settings.from_environment({"BUDDY_ENVIRONMENT": "test"}).uploads is None
    for values in (
        {"BUDDY_UPLOADS_ENABLED": "yes"},
        {"BUDDY_UPLOADS_ENABLED": "true"},
        {"BUDDY_UPLOADS_ENABLED": "true", "BUDDY_UPLOAD_MAX_BYTES": "nan"},
    ):
        with pytest.raises(ConfigurationError):
            Settings.from_environment({"BUDDY_ENVIRONMENT": "test", **values})
    create_app(Settings.from_environment({"BUDDY_ENVIRONMENT": "test"}))
    for changes in (
        {"max_bytes": 0},
        {"expanded_bytes": 999999999},
        {"max_slides": 501},
        {"seconds": float("nan")},
    ):
        with pytest.raises(ValidationError):
            UploadSettings.model_validate({**CONFIG.model_dump(), **changes})


def test_metadata_rejects_owner_override_unsafe_filename_and_wrong_type():
    valid = {
        "filename": "safe.pptx",
        "title": "Title",
        "size_bytes": 1,
        "sha256": "a" * 64,
        "content_type": PPTX_TYPE,
    }
    for changes in (
        {"owner_id": str(uuid4())},
        {"filename": "unsafe.pptm"},
        {"size_bytes": True},
        {"title": "bad\nvalue"},
        {"content_type": "application/zip"},
    ):
        with pytest.raises(ValidationError):
            UploadMetadata.model_validate({**valid, **changes})
