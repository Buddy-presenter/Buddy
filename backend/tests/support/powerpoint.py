"""Synthetic permissioned decks for PowerPoint pipeline acceptance."""

import io
from functools import lru_cache

from pptx import Presentation
from pptx.util import Pt

from buddy_api.analysis.powerpoint import (
    PowerPointOutput,
    SuggestedAction,
    SuggestedSlide,
    powerpoint_manifest,
)
from buddy_api.analysis.provider import ModelSummary
from buddy_api.extraction.runner import extract_pptx


@lru_cache(maxsize=2)
def source(width=960):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(width), Pt(540)
    for heading in ("Quarterly results", "Next steps"):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(Pt(50), Pt(55), Pt(300), Pt(80))
        box.text_frame.text = heading
        body = slide.shapes.add_textbox(Pt(50), Pt(180), Pt(400), Pt(160))
        body.text_frame.text = "Revenue grew 12%. Expand the pilot after review."
        slide.notes_slide.notes_text_frame.text = "Discuss the measured pilot results."
    output = io.BytesIO()
    deck.save(output)
    content = output.getvalue()
    extraction = extract_pptx(io.BytesIO(content), source_revision="issue128-test")
    return content, extraction


def output(extraction, kinds=("highlight_region",)):
    return PowerPointOutput(
        source_format="pptx",
        presentation=ModelSummary(text="Quarterly results and next steps.", reason=None),
        slides=tuple(
            SuggestedSlide(
                number=index + 1,
                summary=ModelSummary(
                    text=slide.text + " Notes: " + (slide.notes or ""), reason=None
                ),
                actions=tuple(
                    SuggestedAction(
                        type=kind,
                        purpose="Emphasize the quarterly result.",
                        object_id=slide.objects[0].id,
                        text="Quarterly result" if kind in {"callout", "text_box"} else "",
                    )
                    for kind in kinds
                ),
            )
            for index, slide in enumerate(extraction.slides)
        ),
    )


def manifest(extraction, deck_id="00000000-0000-0000-0000-000000000128"):
    return powerpoint_manifest(extraction, str(deck_id))
